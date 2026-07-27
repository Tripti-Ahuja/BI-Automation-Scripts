"""
IBM Cognos Analytics Data Extractor  (VDI-safe / IDLE-compatible)
=================================================================
Read-only script — fetches reports, folders, data sources, users from Cognos.
Writes ONLY .xlsx output files to the same folder as this script.
No credentials are saved to disk.

Prerequisites (install via pip):
    pip install requests openpyxl truststore

Security notes:
    - Password is collected via getpass (hidden input) and held in memory only
    - All API calls are read-only; nothing is modified on the server
    - Session is signed out on exit, even if the script crashes
"""
import sys
import os
import re
import ssl
import time
import getpass
import warnings
import xml.etree.ElementTree as ET
from datetime import datetime

try:
    import requests
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print("")
    print("  Missing package: " + str(e))
    print("")
    print("  Install these in your VDI (run in cmd or IDLE terminal):")
    print("    pip install requests openpyxl truststore")
    print("")
    input("  Press Enter to close...")
    sys.exit(1)


# ---------------------------------------------------------------------------
# SSL / TLS — use the OS certificate store (critical for corporate VDI)
# ---------------------------------------------------------------------------
_SSL_VERIFY = True
_SSL_STATUS = "unknown"

try:
    import truststore
    truststore.inject_into_ssl()
    _SSL_STATUS = "truststore (OS cert store)"
except ImportError:
    try:
        import certifi
        os.environ.setdefault("REQUESTS_CA_BUNDLE", certifi.where())
        _SSL_STATUS = "certifi bundle"
    except ImportError:
        _ctx = ssl.create_default_context()
        if _ctx.get_ca_certs():
            _SSL_STATUS = "Python default SSL context"
        else:
            import urllib3
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
            _SSL_VERIFY = False
            _SSL_STATUS = "DISABLED (no cert store found)"


MAX_RETRIES = 3
RETRY_DELAY = 5

REPORT_TYPES = {"report", "reportView", "exploration", "dashboard", "story",
                "interactiveReport", "powerPlayReport", "analysis", "query",
                "jupyterNotebook"}
FOLDER_TYPES = {"folder", "package", "namespaceFolder"}
DATASOURCE_TYPES = {"dataSource", "dataSourceConnection", "dataSourceSignon",
                    "dataModule"}

_ATOM_NS = {
    "atom": "http://www.w3.org/2005/Atom",
    "cm": "http://developer.cognos.com/schemas/bibus/3/",
}


# ---------------------------------------------------------------------------
# URL parsing
# ---------------------------------------------------------------------------
def _parse_url(raw):
    """Extract base URL (server + context path) from a Cognos browser URL."""
    raw = raw.strip().rstrip("/")
    if not raw.startswith("http"):
        raw = "https://" + raw
    for sep in ("?", "#"):
        i = raw.find(sep)
        if i != -1:
            raw = raw[:i].rstrip("/")
    return raw


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------
def api_get(session, url):
    """GET with retry logic. Returns response object or None on failure."""
    for attempt in range(MAX_RETRIES):
        try:
            r = session.get(url, timeout=60)
        except Exception:
            if attempt < MAX_RETRIES - 1:
                time.sleep(RETRY_DELAY * (attempt + 1))
                continue
            return None
        if r.status_code == 200:
            return r
        if r.status_code == 429:
            wait = int(r.headers.get("Retry-After", RETRY_DELAY * (attempt + 1)))
            time.sleep(wait)
            continue
        if attempt < MAX_RETRIES - 1:
            time.sleep(RETRY_DELAY)
            continue
        return None
    return None


def _parse_json(r):
    if r is None:
        return []
    try:
        data = r.json()
    except Exception:
        return []
    if isinstance(data, list):
        return data
    return data.get("content", data.get("data", data.get("value", [])))


def _parse_atom_entry(entry):
    item = {}
    title = entry.find("atom:title", _ATOM_NS) or entry.find("title")
    item["defaultName"] = title.text if title is not None and title.text else ""
    eid = entry.find("atom:id", _ATOM_NS) or entry.find("id")
    raw_id = eid.text if eid is not None and eid.text else ""
    m = re.search(r'storeID\("([^"]+)"\)', raw_id)
    item["id"] = m.group(1) if m else raw_id
    upd = entry.find("atom:updated", _ATOM_NS) or entry.find("updated")
    item["modificationTime"] = upd.text if upd is not None else ""
    pub = entry.find("atom:published", _ATOM_NS) or entry.find("published")
    item["creationTime"] = pub.text if pub is not None else ""

    content_el = entry.find("atom:content", _ATOM_NS) or entry.find("content")
    if content_el is not None:
        for child in content_el:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "objectClass":
                item["type"] = child.text or ""
            elif tag == "defaultDescription":
                item["defaultDescription"] = child.text or ""
            elif tag == "hidden":
                item["hidden"] = (child.text == "true")
            elif tag == "owner":
                for sub in child:
                    if (sub.tag.split("}")[-1] if "}" in sub.tag else sub.tag) == "defaultName":
                        item["owner"] = {"defaultName": sub.text or ""}

    if "type" not in item:
        for cat in list(entry.findall("atom:category", _ATOM_NS)) + list(entry.findall("category")):
            term = cat.get("term", "")
            if term:
                item["type"] = term
                break
        item.setdefault("type", "")
    return item


def _parse_atom(r):
    items = []
    if r is None:
        return items
    try:
        root = ET.fromstring(r.text)
    except Exception:
        return items
    entries = (root.findall("atom:entry", _ATOM_NS)
               or root.findall("entry")
               or root.findall(".//{http://www.w3.org/2005/Atom}entry"))
    for entry in entries:
        it = _parse_atom_entry(entry)
        if it.get("defaultName") or it.get("id"):
            items.append(it)
    return items


# ---------------------------------------------------------------------------
# Authentication
# ---------------------------------------------------------------------------
def login():
    """Authenticate to Cognos Analytics. Returns (session, base_url, endpoints) or None."""
    print("")
    print("  === COGNOS ANALYTICS DATA EXTRACTOR ===")
    print("  (read-only / VDI-safe)")
    print("")
    print("  SSL verification: " + _SSL_STATUS)
    if not _SSL_VERIFY:
        print("")
        print("  !! WARNING: SSL verification is OFF.")
        print("     Install truststore to fix:  pip install truststore")
        if input("     Continue anyway? [y/N]: ").strip().lower() not in ("y", "yes"):
            return None
    print("")

    raw_url = input("  Cognos URL (paste your browser URL): ").strip()
    base_url = _parse_url(raw_url)
    print("  -> Base URL: " + base_url)
    print("")

    namespace = input("  Namespace (e.g. LDAP, CognosAD): ").strip()
    username = input("  Username (try DOMAIN\\user if plain fails): ").strip()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", getpass.GetPassWarning)
        try:
            password = getpass.getpass("  Password: ")
        except Exception:
            password = input("  Password (visible in IDLE): ")

    if not (namespace and username and password):
        print("  All fields are required.")
        return None

    session = requests.Session()
    session.verify = _SSL_VERIFY
    session.headers.update({
        "Accept": "application/json, application/atom+xml, */*",
    })

    json_body = {"parameters": [
        {"name": "CAMNamespace", "value": namespace},
        {"name": "CAMUsername", "value": username},
        {"name": "CAMPassword", "value": password},
    ]}
    form_body = {"CAMNamespace": namespace, "CAMUsername": username, "CAMPassword": password}

    # Modern endpoints first — give a session valid for /v1/objects/* and /api/v1/*.
    # Legacy /rds/auth/logon last — its cam_passport only works for /rds/* and
    # is rejected by modern content endpoints (RDS-ERR-1020).
    attempts = [
        ("PUT  /v1/login (JSON)",        "put",  base_url + "/v1/login",               {"json": json_body}),
        ("POST /v1/login (JSON)",        "post", base_url + "/v1/login",               {"json": json_body}),
        ("PUT  /api/v1/session (JSON)",  "put",  base_url + "/api/v1/session",         {"json": json_body}),
        ("POST /api/v1/session (JSON)",  "post", base_url + "/api/v1/session",         {"json": json_body}),
        ("POST /v1/disp/rds/auth/logon", "post", base_url + "/v1/disp/rds/auth/logon", {"data": form_body}),
        ("POST /bi/v1/disp/rds/auth/logon", "post", base_url + "/bi/v1/disp/rds/auth/logon", {"data": form_body}),
    ]

    print("")
    print("  Signing in...")
    success_label = None
    last_status = ""
    for label, method_name, url, kwargs in attempts:
        try:
            r = getattr(session, method_name)(url, timeout=30, **kwargs)
        except Exception:
            continue
        has_passport = "cam_passport" in session.cookies.get_dict()
        if r.status_code in (200, 201) or has_passport:
            success_label = label
            # Capture XSRF token (header first, then cookie)
            for h in ("X-XSRF-Token", "XSRF-Token", "X-CAM-XSRF"):
                if r.headers.get(h):
                    session.headers["X-XSRF-Token"] = r.headers[h]
                    break
            if "X-XSRF-Token" not in session.headers:
                for c in ("XSRF-TOKEN", "X-XSRF-TOKEN", "X-CAM-XSRF"):
                    if session.cookies.get(c):
                        session.headers["X-XSRF-Token"] = session.cookies.get(c)
                        break
            print("  Success via " + label)
            break
        last_status = "HTTP " + str(r.status_code) + " on " + label

    del password

    if not success_label:
        print("")
        print("  LOGIN FAILED — last: " + last_status)
        print("  - Verify namespace matches the Cognos login dropdown exactly")
        print("  - Try username formats: just user / user@domain.com / DOMAIN\\user")
        print("  - Confirm the password is correct (same as browser)")
        return None

    # Discover which content endpoint accepts this session
    print("")
    print("  Detecting content endpoint...")
    candidates = [
        # Modern object API (preferred when /v1/login succeeded)
        ("rest", base_url + "/v1/objects/.public_folders/items",
                 base_url + "/v1/objects/{id}/items"),
        ("rest", base_url + "/v1/objects/.my_folders/items",
                 base_url + "/v1/objects/{id}/items"),
        # Public REST API
        ("rest", base_url + "/api/v1/folders/.public_folders/items",
                 base_url + "/api/v1/folders/{id}/items"),
        ("rest", base_url + "/api/v1/folders/.my_folders/items",
                 base_url + "/api/v1/folders/{id}/items"),
        ("rest", base_url + "/api/v1/content",
                 base_url + "/api/v1/content/{id}/items"),
        # Dispatch ATOM (legacy fallback)
        ("atom", base_url + "/v1/disp/rds/atom/content",
                 base_url + '/v1/disp/rds/atom/content/storeID("{id}")/item'),
    ]

    endpoints = None
    for mode, root_url, children in candidates:
        try:
            r = session.get(root_url, timeout=20)
        except Exception:
            continue
        if r.status_code != 200:
            continue
        if mode == "rest":
            try:
                data = r.json()
                items = (data if isinstance(data, list)
                         else data.get("content", data.get("data", data.get("value", []))))
                if not isinstance(items, list):
                    continue
            except Exception:
                continue
        else:
            if "<entry" not in r.text and "<feed" not in r.text:
                continue
        endpoints = {"mode": mode, "root": root_url, "children": children}
        print("  -> Using " + mode.upper() + " :  " + root_url)
        break

    if endpoints is None:
        print("")
        print("  ERROR: Authenticated, but no content endpoint accepts this session.")
        print("  Your Cognos server likely has the REST API disabled and the legacy")
        print("  dispatch session doesn't have content-read permission.")
        print("  Ask your Cognos admin to enable the REST API on this dispatcher.")
        return None

    print("")
    return session, base_url, endpoints


# ---------------------------------------------------------------------------
# Content walker
# ---------------------------------------------------------------------------
def _fetch_items(session, endpoints, parent_id=None):
    if parent_id is None:
        url = endpoints["root"]
    else:
        url = endpoints["children"].replace("{id}", parent_id)
    r = api_get(session, url)
    if r is None:
        return []
    if endpoints["mode"] == "rest":
        return _parse_json(r)
    return _parse_atom(r)


def _walk(session, endpoints, parent_id=None, parent_path="",
          collect_types=None, results=None, depth=0):
    if results is None:
        results = []
    if depth > 20:
        return results

    items = _fetch_items(session, endpoints, parent_id)
    if depth == 0:
        print("  [" + str(len(items)) + " top-level items]")

    for item in items:
        item_type = item.get("type", "")
        item_name = item.get("defaultName", item.get("name", ""))
        full_path = (parent_path + " / " + item_name) if parent_path else item_name

        if collect_types is None or item_type in collect_types:
            item["_full_path"] = full_path
            results.append(item)

        if item_type in FOLDER_TYPES:
            item_id = item.get("id", "")
            if item_id:
                _walk(session, endpoints, item_id, full_path,
                      collect_types, results, depth + 1)

    return results


# ---------------------------------------------------------------------------
# Row helpers
# ---------------------------------------------------------------------------
def _get_owner(item):
    owner = item.get("owner", {})
    if isinstance(owner, dict):
        return owner.get("defaultName", owner.get("name", owner.get("id", "")))
    return str(owner) if owner else ""


def _fmt_date(val):
    return str(val).replace("T", " ").replace("Z", "") if val else ""


def _content_row(item):
    return {
        "Name": item.get("defaultName", item.get("name", "")),
        "Full Path": item.get("_full_path", ""),
        "Type": item.get("type", ""),
        "ID": item.get("id", ""),
        "Owner": _get_owner(item),
        "Created": _fmt_date(item.get("creationTime", "")),
        "Modified": _fmt_date(item.get("modificationTime", "")),
        "Description": item.get("defaultDescription", item.get("description", "")),
        "Hidden": item.get("hidden", False),
    }


# ===================================================================
#  OPTION 1:  REPORTS & DASHBOARDS
# ===================================================================
def fetch_reports(session, base_url, endpoints):
    print("  [Scanning content store for reports & dashboards...]")
    items = _walk(session, endpoints, collect_types=REPORT_TYPES)
    rows = [_content_row(i) for i in items]
    print("  Done - " + str(len(rows)) + " reports/dashboards")
    print("")
    return rows


# ===================================================================
#  OPTION 2:  FOLDERS & PACKAGES
# ===================================================================
def fetch_folders(session, base_url, endpoints):
    print("  [Scanning content store for folders & packages...]")
    items = _walk(session, endpoints, collect_types=FOLDER_TYPES)
    rows = [_content_row(i) for i in items]
    print("  Done - " + str(len(rows)) + " folders/packages")
    print("")
    return rows


# ===================================================================
#  OPTION 3:  DATA SOURCES
# ===================================================================
def fetch_datasources(session, base_url, endpoints):
    print("  [Scanning content store for data sources...]")
    items = _walk(session, endpoints, collect_types=DATASOURCE_TYPES)
    rows = [_content_row(i) for i in items]
    print("  Done - " + str(len(rows)) + " data sources")
    print("")
    return rows


# ===================================================================
#  OPTION 4:  USERS, GROUPS & ROLES
# ===================================================================
def _fetch_namespace_endpoint(session, base_url, endpoints, endpoint_path):
    """Fetch users/groups/roles using the right base path for this mode."""
    if endpoints["mode"] == "rest":
        if "/v1/objects" in endpoints["root"]:
            url = base_url + "/v1/" + endpoint_path
        else:
            url = base_url + "/api/v1/" + endpoint_path
    else:
        url = base_url + "/v1/disp/rds/atom/" + endpoint_path
    r = api_get(session, url)
    if r is None:
        return []
    try:
        return _parse_json(r)
    except Exception:
        return _parse_atom(r)


def fetch_users(session, base_url, endpoints):
    print("  [Getting users, groups & roles...]")

    print("    Fetching users...")
    users = _fetch_namespace_endpoint(session, base_url, endpoints, "users")
    user_rows = [{
        "Display Name": u.get("defaultName", u.get("name", "")),
        "User ID": u.get("id", ""),
        "Email": u.get("email", ""),
        "User Name": u.get("userName", u.get("searchPath", "")),
        "Type": u.get("type", ""),
        "Active": u.get("active", ""),
        "Created": _fmt_date(u.get("creationTime", "")),
        "Modified": _fmt_date(u.get("modificationTime", "")),
    } for u in users]
    print("    Users:  " + str(len(user_rows)))

    print("    Fetching groups...")
    groups = _fetch_namespace_endpoint(session, base_url, endpoints, "groups")
    group_rows = []
    for g in groups:
        members = g.get("members", [])
        members_str = ", ".join(m.get("defaultName", m.get("name", ""))
                                for m in members) if isinstance(members, list) else ""
        group_rows.append({
            "Group Name": g.get("defaultName", g.get("name", "")),
            "Group ID": g.get("id", ""),
            "# Members": len(members) if isinstance(members, list) else "",
            "Members": members_str,
            "Type": g.get("type", ""),
            "Created": _fmt_date(g.get("creationTime", "")),
            "Modified": _fmt_date(g.get("modificationTime", "")),
        })
    print("    Groups: " + str(len(group_rows)))

    print("    Fetching roles...")
    roles = _fetch_namespace_endpoint(session, base_url, endpoints, "roles")
    role_rows = []
    for role in roles:
        members = role.get("members", [])
        members_str = ", ".join(m.get("defaultName", m.get("name", ""))
                                for m in members) if isinstance(members, list) else ""
        role_rows.append({
            "Role Name": role.get("defaultName", role.get("name", "")),
            "Role ID": role.get("id", ""),
            "# Members": len(members) if isinstance(members, list) else "",
            "Members": members_str,
            "Type": role.get("type", ""),
            "Created": _fmt_date(role.get("creationTime", "")),
            "Modified": _fmt_date(role.get("modificationTime", "")),
        })
    print("    Roles:  " + str(len(role_rows)))
    print("")
    return user_rows, group_rows, role_rows


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
_HDR_FONT = Font(bold=True, color="FFFFFF")
_HDR_FILL = PatternFill("solid", fgColor="2F5496")
_HDR_ALIGN = Alignment(horizontal="center")


def _write_sheet(ws, rows):
    if not rows:
        ws["A1"] = "No data"
        return
    heads = list(rows[0].keys())
    for ci, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font, c.fill, c.alignment = _HDR_FONT, _HDR_FILL, _HDR_ALIGN
    for ri, row in enumerate(rows, 2):
        for ci, h in enumerate(heads, 1):
            v = row.get(h, "")
            ws.cell(row=ri, column=ci,
                    value=str(v) if isinstance(v, (dict, list)) else v)
    for ci, h in enumerate(heads, 1):
        sample_lens = [len(str(row.get(h, ""))) for row in rows[:200]]
        max_len = max(len(h), max(sample_lens)) if sample_lens else len(h)
        ws.column_dimensions[get_column_letter(ci)].width = min(max(max_len + 2, 14), 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _out_path(name):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "Cognos_" + name + "_" + ts + ".xlsx")


def _save_multi(sheets, filename):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        _write_sheet(ws, rows)
    path = _out_path(filename)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Interactive CLI
# ---------------------------------------------------------------------------
def main():
    result = login()
    if not result:
        input("\n  Press Enter to close...")
        return
    session, base_url, endpoints = result

    try:
        while True:
            print("")
            print("  OPTIONS")
            print("  " + "-" * 45)
            print("  1. Reports & Dashboards")
            print("  2. Folders & Packages")
            print("  3. Data Sources")
            print("  4. Users, Groups & Roles")
            print("  5. Export ALL (everything in one file)")
            print("  0. Exit")

            ch = input("\n  Pick (0-5): ").strip()

            if ch == "0":
                break

            elif ch == "1":
                rows = fetch_reports(session, base_url, endpoints)
                path = _save_multi({"Reports & Dashboards": rows}, "Reports_Dashboards")
                print("  Saved -> " + path + "  (" + str(len(rows)) + " rows)")
                input("\n  Press Enter to go back...")

            elif ch == "2":
                rows = fetch_folders(session, base_url, endpoints)
                path = _save_multi({"Folders & Packages": rows}, "Folders_Packages")
                print("  Saved -> " + path + "  (" + str(len(rows)) + " rows)")
                input("\n  Press Enter to go back...")

            elif ch == "3":
                rows = fetch_datasources(session, base_url, endpoints)
                path = _save_multi({"Data Sources": rows}, "DataSources")
                print("  Saved -> " + path + "  (" + str(len(rows)) + " rows)")
                input("\n  Press Enter to go back...")

            elif ch == "4":
                user_rows, group_rows, role_rows = fetch_users(session, base_url, endpoints)
                sheets = {"Users": user_rows, "Groups": group_rows, "Roles": role_rows}
                path = _save_multi(sheets, "Users_Groups_Roles")
                print("  Saved -> " + path)
                print("    Users:  " + str(len(user_rows)))
                print("    Groups: " + str(len(group_rows)))
                print("    Roles:  " + str(len(role_rows)))
                input("\n  Press Enter to go back...")

            elif ch == "5":
                print("")
                print("  Fetching everything...")
                print("")
                print("  --- Reports & Dashboards ---")
                reports = fetch_reports(session, base_url, endpoints)
                print("  --- Folders & Packages ---")
                folders = fetch_folders(session, base_url, endpoints)
                print("  --- Data Sources ---")
                datasources = fetch_datasources(session, base_url, endpoints)
                print("  --- Users, Groups & Roles ---")
                user_rows, group_rows, role_rows = fetch_users(session, base_url, endpoints)

                sheets = {
                    "Reports & Dashboards": reports,
                    "Folders & Packages": folders,
                    "Data Sources": datasources,
                    "Users": user_rows,
                    "Groups": group_rows,
                    "Roles": role_rows,
                }
                path = _save_multi(sheets, "ALL")
                total = sum(len(v) for v in sheets.values())
                print("  Exported " + str(total) + " total rows -> " + path)
                print("    Reports & Dashboards: " + str(len(reports)))
                print("    Folders & Packages:   " + str(len(folders)))
                print("    Data Sources:         " + str(len(datasources)))
                print("    Users:                " + str(len(user_rows)))
                print("    Groups:               " + str(len(group_rows)))
                print("    Roles:                " + str(len(role_rows)))
                print("")
                print("  NOTE: file may contain sensitive data (server names, usernames)")
                input("\n  Press Enter to go back...")

            else:
                print("  Invalid choice.")

    finally:
        for path in ("/v1/login", "/api/v1/session", "/v1/disp/rds/auth/logoff"):
            try:
                session.delete(base_url + path, timeout=10)
            except Exception:
                pass
        print("\n  Signed out from Cognos Analytics.")

    print("  Bye!\n")


if __name__ == "__main__":
    main()
